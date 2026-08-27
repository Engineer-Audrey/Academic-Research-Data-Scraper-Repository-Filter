import time
import keyboard
import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from openpyxl.styles import Alignment, PatternFill, Font, Border, Side

# Define the keywords to search in the Mitacs catalog
KEYWORDS_TO_SEARCH = [
    "Machine Learning",
    "Artificial Intelligence",
    "Java",
    "Software Engineering",
    "Database",
    "Data"
]

# Target cities to match projects against
TARGET_CITIES = [
    "Montreal",   # Québec (Light Pastel Blue)
    "Calgary",    # Alberta (Light Pastel Green)
    "Edmonton",   # Alberta
    "Ottawa",     # Ontario (Light Pastel Pink)
    "Kingston",   # Ontario
    "London"      # Ontario
]

MAX_PER_CITY_PER_KEYWORD = 10

# Initialize Chrome Driver
options = webdriver.ChromeOptions()
driver = webdriver.Chrome(options=options)

try:
    print("Opening Mitacs Globalink Portal...")
    driver.get("https://globalink.mitacs.ca/#/student/application/projects")
    time.sleep(8) 

    projects_by_keyword = {}
    seen_ids = set()

    # -------------------------------------------------------------
    # KEYWORD SEARCH LOOP
    # -------------------------------------------------------------
    for kw in KEYWORDS_TO_SEARCH:
        print(f"\n=======================================================")
        print(f"---> SEARCHING FOR KEYWORD: '{kw}' (MAX {MAX_PER_CITY_PER_KEYWORD} PER CITY)")
        print(f"=======================================================")
        print("💡 TIP: Press the 'RIGHT ARROW (->)' key at any time to skip to the next keyword.")

        current_kw_projects = []

        try:
            # A. Enter keyword in Keyword search box
            search_box = driver.find_element(By.CSS_SELECTOR, "input.p-inputtext")
            search_box.clear()
            search_box.send_keys(kw)
            time.sleep(1)

            # B. Select Language filter as 'English'
            try:
                lang_dropdowns = driver.find_elements(By.XPATH, "//div[contains(., 'Language') and .//p-dropdown]")
                if lang_dropdowns:
                    lang_dd = lang_dropdowns[0].find_element(By.TAG_NAME, "p-dropdown")
                    lang_dd.click()
                    time.sleep(1)
                    
                    option_eng = driver.find_element(By.XPATH, "//li[contains(@class, 'p-dropdown-item') and contains(., 'English')]")
                    option_eng.click()
                    time.sleep(1)
            except Exception as e_lang:
                print(f"Note: Adjusting Language filter to English: {e_lang}")

            # C. Select 'Computer Science' under academic discipline/background
            try:
                dropdowns = driver.find_elements(By.CSS_SELECTOR, "p-dropdown")
                for dropdown in dropdowns[-2:]:
                    dropdown.click()
                    time.sleep(1)
                    
                    filter_input = driver.find_element(By.CSS_SELECTOR, "input.p-dropdown-filter")
                    filter_input.clear()
                    filter_input.send_keys("Computer Science")
                    time.sleep(1)

                    option_cs = driver.find_element(By.XPATH, "//li[contains(@class, 'p-dropdown-item') and contains(., 'Computer Science')]")
                    option_cs.click()
                    time.sleep(1)
            except Exception:
                pass

            # D. Submit search query
            try:
                btn_search = driver.find_element(By.XPATH, "//button[contains(., 'Search and Filter')]")
                btn_search.click()
            except Exception:
                search_box.send_keys(Keys.ENTER)

            time.sleep(6) 

        except Exception as e:
            print(f"Error applying keyword '{kw}': {e}")
            continue

        city_counts = {city.lower(): 0 for city in TARGET_CITIES}
        page_num = 1
        previous_active_page = ""
        no_new_data_counter = 0

        while True:
            # 1. Manual override check (Right Arrow key)
            if keyboard.is_pressed('right'):
                print(f"\n[RIGHT ARROW DETECTED] Skipping keyword '{kw}' as requested...")
                time.sleep(1)
                break

            # 2. Early stopping rule if 3 target cities reach the maximum quota
            cities_at_max = sum(1 for count in city_counts.values() if count >= MAX_PER_CITY_PER_KEYWORD)
            if cities_at_max >= 3:
                print(f"Quota reached (10 projects) in at least 3 target cities! Moving to next keyword...")
                break

            # 3. VERIFICATION 1: Active page reading from PrimeNG paginator
            try:
                active_page_elem = driver.find_element(By.CSS_SELECTOR, ".p-paginator-page.p-highlight, .p-paginator-page.p-state-active")
                current_active_page = active_page_elem.text.strip()
            except Exception:
                current_active_page = str(page_num)

            if current_active_page == previous_active_page and page_num > 1:
                print(f"End of available pagination automatically reached at page {current_active_page} for '{kw}'.")
                break
            
            previous_active_page = current_active_page

            print(f"-> Scraping real page {current_active_page} for '{kw}'... Current counts: {city_counts}")

            cards = driver.find_elements(By.XPATH, "//div[contains(., 'Project ID') and .//button[contains(text(), 'View Detail')]]")
            if not cards:
                cards = driver.find_elements(By.XPATH, "//*[contains(text(), 'Project ID')]/ancestor::div[2]")

            new_projects_in_this_page = 0

            for card in cards:
                try:
                    raw_text = card.text.strip()
                    lines = [line.strip() for line in raw_text.split('\n') if line.strip()]

                    if raw_text.count("Project ID") > 1:
                        continue

                    project_id = ""
                    title = ""
                    university = ""
                    campus = ""
                    province = ""
                    city = ""
                    language = ""
                    start_date = ""
                    description = ""

                    for i, line in enumerate(lines):
                        if "Project ID" in line:
                            project_id = line.replace("Project ID", "").strip()
                            if i + 1 < len(lines):
                                title = lines[i + 1]
                        elif "Faculty University:" in line:
                            university = line.replace("Faculty University:", "").strip()
                        elif "Faculty Campus:" in line:
                            campus = line.replace("Faculty Campus:", "").strip()
                        elif "Faculty Province:" in line:
                            province = line.replace("Faculty Province:", "").strip()
                        elif "Project Location:" in line:
                            city = line.replace("Project Location:", "").strip()
                        elif "Language:" in line:
                            language = line.replace("Language:", "").strip()
                        elif "Preferred start" in line:
                            start_date = line.replace("Preferred start date:", "").replace("Preferred start:", "").strip()

                    # Filter out non-English projects
                    if language and "english" not in language.lower():
                        continue

                    # Enhanced Start Date Filter (Handles YYYY-05-DD, "May", "As soon as possible", or empty dates)
                    if start_date:
                        s_lower = start_date.lower()
                        is_may = ("may" in s_lower) or ("-05-" in s_lower) or ("/05/" in s_lower) or s_lower.endswith("-05")
                        is_asap = ("as soon as possible" in s_lower) or ("asp" in s_lower)
                        if not (is_may or is_asap):
                            continue

                    candidates = [l for l in lines if len(l) > 70 and not any(k in l for k in ["Faculty", "Project Location", "Language", "Preferred start", "Project ID"])]
                    if candidates:
                        description = candidates[0]

                    matched_city = None
                    for target in TARGET_CITIES:
                        if target.lower() in city.lower() or target.lower() in campus.lower():
                            matched_city = target.lower()
                            break

                    if project_id and project_id not in seen_ids and matched_city:
                        if city_counts[matched_city] < MAX_PER_CITY_PER_KEYWORD:
                            seen_ids.add(project_id)
                            city_counts[matched_city] += 1
                            new_projects_in_this_page += 1
                            
                            current_kw_projects.append({
                                "Keyword Used": kw,
                                "Page": current_active_page,
                                "Start Date": start_date if start_date else "As soon as possible",
                                "Project ID": project_id,
                                "Title": title,
                                "University": university,
                                "Campus": campus,
                                "Province": province,
                                "City": city,
                                "Language": language,
                                "Description": description
                            })
                            print(f"  + Registered Project [{project_id}] (Page {current_active_page} | Start: {start_date or 'N/A'}) in {city} ({city_counts[matched_city]}/{MAX_PER_CITY_PER_KEYWORD} for '{kw}')")

                except Exception:
                    continue

            # VERIFICATION 2: Check for empty/repeated consecutive pages
            if new_projects_in_this_page == 0:
                no_new_data_counter += 1
                if no_new_data_counter >= 2 and page_num > 5:
                    print(f"No new projects detected for 2 consecutive pages. Stopping pagination for '{kw}'.")
                    break
            else:
                no_new_data_counter = 0

            # VERIFICATION 3: Click next page or break if disabled
            try:
                next_btn = driver.find_element(By.XPATH, "//button[contains(@class, 'p-paginator-next')]")
                btn_class = next_btn.get_attribute("class") or ""
                is_disabled = next_btn.get_attribute("disabled") or next_btn.get_attribute("aria-disabled")

                if "p-disabled" in btn_class or is_disabled == "true" or is_disabled is not None or not next_btn.is_enabled():
                    print(f"'Next' button is disabled. Pagination complete for '{kw}'.")
                    break
                else:
                    next_btn.click()
                    page_num += 1
                    time.sleep(4)
            except Exception:
                print(f"No more pagination controls found for '{kw}'.")
                break

        if current_kw_projects:
            projects_by_keyword[kw] = current_kw_projects

    # -------------------------------------------------------------
    # EXPORT TO EXCEL WITH CUSTOM STYLES & TABBED SHEETS
    # -------------------------------------------------------------
    if projects_by_keyword:
        excel_filename = "mitacs_projects.xlsx"
        
        with pd.ExcelWriter(excel_filename, engine='openpyxl') as writer:
            
            # Fills by province (Pastel Colors)
            fill_quebec = PatternFill(start_color="D0E0E3", end_color="D0E0E3", fill_type="solid")   # Pastel Light Blue
            fill_alberta = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Pastel Light Green
            fill_ontario = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")  # Pastel Light Pink
            fill_other = PatternFill(start_color="F3F3F3", end_color="F3F3F3", fill_type="solid")
            
            # Soft Pastel Highlight for Start Date (Pastel Light Yellow/Cream)
            fill_start_date = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

            # Header styles
            fill_header = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")  # Light Gray
            font_header = Font(name="Calibri", size=14, bold=True, color="000000")

            # Thin borders
            thin_side = Side(border_style="thin", color="D3D3D3")
            cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

            wrap_alignment = Alignment(wrap_text=True, vertical='top')
            header_alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

            col_widths = {
                'A': 22,  # Keyword Used
                'B': 10,  # Page
                'C': 22,  # Start Date
                'D': 14,  # Project ID
                'E': 32,  # Title
                'F': 28,  # University
                'G': 20,  # Campus
                'H': 16,  # Province
                'I': 22,  # City
                'J': 14,  # Language
                'K': 70   # Description
            }

            total_projects_count = 0

            for kw_name, data_list in projects_by_keyword.items():
                df = pd.DataFrame(data_list)
                df.sort_values(by=["Province", "City"], inplace=True)
                
                sheet_title = kw_name[:30].replace(":", "").replace("/", "")
                df.to_excel(writer, index=False, sheet_name=sheet_title)
                
                worksheet = writer.sheets[sheet_title]

                # 1. Apply column widths
                for col_letter, width in col_widths.items():
                    worksheet.column_dimensions[col_letter].width = width

                # 2. Format Header Row (Row 1)
                for cell in worksheet[1]:
                    cell.font = font_header
                    cell.fill = fill_header
                    cell.alignment = header_alignment
                    cell.border = cell_border

                # 3. Format Data Rows (Row 2 onwards)
                for row in worksheet.iter_rows(min_row=2, max_col=11, max_row=len(data_list)+1):
                    prov_val = str(row[7].value).lower() if row[7].value else ""  # Column H is Province
                    
                    if "qu&eacute;bec" in prov_val or "quebec" in prov_val:
                        prov_fill = fill_quebec
                    elif "alberta" in prov_val:
                        prov_fill = fill_alberta
                    elif "ontario" in prov_val:
                        prov_fill = fill_ontario
                    else:
                        prov_fill = fill_other

                    for col_idx, cell in enumerate(row, start=1):
                        cell.alignment = wrap_alignment
                        cell.border = cell_border
                        
                        # Column A ("Keyword Used"): Plain fill (None)
                        if col_idx == 1:
                            cell.fill = PatternFill(fill_type=None)
                        # Column C ("Start Date"): Soft Highlight Pastel Yellow/Cream
                        elif col_idx == 3:
                            cell.fill = fill_start_date
                        else:
                            cell.fill = prov_fill

                total_projects_count += len(data_list)

        print(f"\nSUCCESS! Saved {total_projects_count} projects in '{excel_filename}'.")
    else:
        print("No projects found matching the selected criteria.")

finally:
    # driver.quit()
    pass